"""جودة المخرج: التنقية من الملاحظات الإنجليزية والعلامات المائية، واللغة البشرية، والعرض والتنزيل."""
import io
import json
import sys
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "runtime"))

from test_panel import fake_claude, panel  # noqa: E402,F401

RAW = ("# العنوان\n\nد. ماجد بوشليبي\n\n## المقدمة\n\n===BEGIN_TEXT===\n"
       "تتناول الدراسة​ الأثر [FACT] الثقافي للمهرجانات [NEEDS-EVIDENCE: exact figure]. "
       "[AUTHOR — بانتظار الاعتماد: الأطروحة بين (أ) و(ب)] و[AUTHOR — pending approval] "
       "ومفهوم رأس المال الثقافي (cultural capital) [@SRC-001] ⟨يحتاج توثيقاً⟩.\n===END_TEXT===\n\nNOTES: لا شيء")


def test_clean_translates_notes_and_strips_watermarks():
    from rkpos import sanitize as S
    text, rep = S.clean(RAW)
    assert "​" not in text and " " not in text and "===" not in text and "@SRC" not in text
    assert "[" not in text and "⟨" not in text
    assert "(يحتاج توثيقاً)" in text and "(قرار المؤلف: بانتظار الاعتماد: الأطروحة بين (أ) و(ب))" in text
    assert "(قرار المؤلف: بانتظار الاعتماد)" in text           # عبارة إنجليزية معروفة تُرجمت
    assert "exact figure" not in text and rep["english_removed"] == ["exact figure"]
    assert "(cultural capital)" in text                         # المصطلح بعد مقابله العربي مسموح
    assert "ملاحظات الوكيل:" in text and "NOTES" not in text
    assert "(واقعة)" not in text                                # وسم التصنيف يُحذف من نسخة القراءة
    tagged, _ = S.clean(RAW, "tagged")
    assert "(واقعة)" in tagged and "[FACT]" not in tagged


def test_source_keeps_canonical_tags_but_drops_hidden_chars():
    from rkpos import sanitize as S
    out = S.source("نص‏ [FACT] و﻿[NEEDS-EVIDENCE]")
    assert out == "نص [FACT] و[NEEDS-EVIDENCE]"


def test_humanlang_report_flags_machine_tells():
    from rkpos import sanitize as S
    body = ("تجدر الإشارة إلى أن المهرجان يلعب دوراً محورياً — وفي الختام — كما أن [TODO: check numbers] ثم "
            + "تقرير منضبط عن الأثر الثقافي " * 10)
    r = S.humanlang_report(body)
    assert {"stock_phrases", "em_dash", "latin_notes"} <= set(r["flags"])
    assert "[TODO: check numbers]" in r["hits"]["latin_notes"]
    clean = S.humanlang_report("تتناول الدراسة أثر المهرجانات في الممارسة الثقافية (cultural practice) " * 20)
    assert clean["ok"], clean


def test_word_file_has_no_hidden_or_visible_watermarks():
    from docx import Document
    from rkpos.export import to_docx
    data = to_docx(RAW, "العنوان", author="د. ماجد بوشليبي")
    z = zipfile.ZipFile(io.BytesIO(data))
    assert "docProps/thumbnail.jpeg" not in z.namelist()
    blob = b"".join(z.read(n) for n in z.namelist())
    assert b"python-docx" not in blob and b"Macintosh" not in blob and b"w:rsid" not in blob
    core = z.read("docProps/core.xml").decode()
    assert "د. ماجد بوشليبي" in core
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert [p.text for p in doc.paragraphs].count("العنوان") == 1     # لا تكرار للعنوان
    assert "[" not in text and "​" not in text and "pending" not in text and "===" not in text
    assert "(يحتاج توثيقاً)" in text


def test_word_run_properties_follow_schema_order():
    """يرفض Word الملف حين تخالف خصائص المقطع ترتيب المخطط (rFonts ثم b ثم bCs … ثم sz ثم szCs … ثم rtl)."""
    from rkpos.export import to_docx
    xml = zipfile.ZipFile(io.BytesIO(to_docx("**غامق** وعادي", "ع"))).read("word/document.xml").decode()
    import re
    order = ["rFonts", "b", "bCs", "color", "sz", "szCs", "rtl"]
    for rpr in re.findall(r"<w:rPr>(.*?)</w:rPr>", xml):
        tags = [t for t in re.findall(r"<w:(\w+)", rpr) if t in order]
        assert tags == sorted(tags, key=order.index), tags


def test_panel_reading_view_and_word_download(panel, tmp_path):
    code, r = panel("new_project", {"title": "دراسة", "type": "journal_article"})
    pid = r["data"]["project_id"]
    f = panel.root / "projects" / pid / "manuscript/book_full.md"
    f.write_text(RAW, encoding="utf-8")
    code, r = panel("file", q="?" + urllib.parse.urlencode({"path": f"projects/{pid}/manuscript/book_full.md", "view": "reading"}))
    assert code == 200 and "[" not in r["data"]["text"] and r["data"]["sanitized"]["english_removed"] == ["exact figure"]
    st = json.loads((tmp_path / "panel.json").read_text())
    q = urllib.parse.urlencode({"path": f"projects/{pid}/manuscript/book_full.md", "title": "دراسة", "clean": "1"})
    req = urllib.request.Request(f"http://127.0.0.1:{st['port']}/api/export?{q}", headers={"X-Rkpos-Token": st["token"]})
    with urllib.request.urlopen(req, timeout=60) as resp:
        assert resp.status == 200
        assert "filename*=UTF-8''" in resp.headers["Content-Disposition"]
        data = resp.read()
    assert data[:2] == b"PK" and b"python-docx" not in data


def test_panel_download_ticket_is_single_use(panel, tmp_path):
    code, r = panel("new_project", {"title": "دراسة التذكرة", "type": "journal_article"})
    pid = r["data"]["project_id"]
    (panel.root / "projects" / pid / "manuscript/book_full.md").write_text(RAW, encoding="utf-8")
    code, r = panel("download_ticket", {"kind": "word", "path": f"projects/{pid}/manuscript/book_full.md", "title": "العمل كاملاً"})
    assert code == 200 and r["data"]["name"] == "دراسة التذكرة.docx"     # العنوان الحقيقي بدل التسمية العامة
    st = json.loads((tmp_path / "panel.json").read_text())
    url = f"http://127.0.0.1:{st['port']}/api/dl?ticket={r['data']['ticket']}"
    with urllib.request.urlopen(url, timeout=30) as resp:                 # بلا رمز: التذكرة نفسها هي الإذن
        assert resp.read()[:2] == b"PK"
    try:
        urllib.request.urlopen(url, timeout=30)
        assert False, "التذكرة تُستعمل مرة واحدة"
    except urllib.error.HTTPError as e:
        assert e.code == 410
    code, r = panel("download_ticket", {"kind": "word", "path": "../pyproject.toml"})
    assert code == 400


SOURCES = [{"Source_ID": "SRC-000001", "Title": "The Forms of Capital", "Author": ["Bourdieu, Pierre"], "Year": 1986,
            "Publisher": "Greenwood", "Verification_Status": "VERIFIED"},
           {"Source_ID": "SRC-000002", "Title": "الاقتصاد الثقافي", "Author": ["الجابري، محمد"], "Year": 2010,
            "Journal": "مجلة الثقافة", "Volume": 3, "Issue": 2, "Verification_Status": "PARTIAL"}]
CITED = ("رأس المال الثقافي [@SRC-000001, p. 45] يتراكم [@SRC-000001] ثم (الجابري، 2010، ص 12) "
         "و(Throsby, 2001) و(أ) و(ب) ومنذ (2019) [@SRC-000009] ثم [@SRC-000001].")


def test_citations_become_numbered_notes_in_text_order():
    from rkpos import footnotes as F
    body, notes = F.collect(CITED, SOURCES)
    assert F.to_superscript(body) == "رأس المال الثقافي¹ يتراكم² ثم³ و⁴ و(أ) و(ب) ومنذ (2019)⁵ ثم⁶."
    assert notes[0] == "Bourdieu, Pierre، The Forms of Capital، Greenwood، 1986، ص 45."
    assert notes[1] == "المرجع نفسه."
    assert notes[2].startswith("الجابري، محمد، «الاقتصاد الثقافي»، مجلة الثقافة، مج 3، ع 2، 2010 (لم يُتحقق منه بعد)")
    assert notes[3] == "Throsby, 2001."                        # غير مسجل: يُنقل الاستشهاد نفسه بلا تلفيق
    assert "SRC-000009" in notes[4] and "يحتاج توثيقاً" in notes[4]
    assert notes[5] == "Bourdieu، The Forms of Capital."        # الصيغة المختصرة لما سبق ذكره


def test_word_has_page_footnotes_and_red_brackets():
    from rkpos.export import to_docx
    data = to_docx("## المقدمة (الإطار)\n\n" + CITED, "عنوان", sources=SOURCES)
    z = zipfile.ZipFile(io.BytesIO(data))
    fn = z.read("word/footnotes.xml").decode()
    doc = z.read("word/document.xml").decode()
    assert fn.count("<w:footnoteRef/>") == 6 and 'w:type="separator"' in fn
    assert doc.count("<w:footnoteReference") == 6 and "@SRC" not in doc and "(Throsby" not in doc
    assert "footnotes.xml" in z.read("word/_rels/document.xml.rels").decode()
    assert "footnotes+xml" in z.read("[Content_Types].xml").decode()
    import re
    red = re.findall(r'<w:color w:val="FF0000"/>.*?<w:t[^>]*>([^<]*)</w:t>', doc)
    assert "(أ)" in red and "(ب)" in red and "(الإطار)" in red and "(2019)" in red
    assert not any("رأس المال" in t for t in red)                 # ما خارج الأقواس بلونه


def test_reader_view_shows_superscripts_and_notes(panel):
    code, r = panel("new_project", {"title": "دراسة الحواشي", "type": "journal_article"})
    pid = r["data"]["project_id"]
    (panel.root / "projects" / pid / "research/sources.jsonl").write_text(
        "\n".join(json.dumps(s, ensure_ascii=False) for s in SOURCES) + "\n", encoding="utf-8")
    (panel.root / "projects" / pid / "manuscript/book_full.md").write_text(CITED, encoding="utf-8")
    code, r = panel("file", q="?" + urllib.parse.urlencode({"path": f"projects/{pid}/manuscript/book_full.md", "view": "reading"}))
    assert "¹" in r["data"]["text"] and "@SRC" not in r["data"]["text"]
    assert r["data"]["notes"][0].startswith("Bourdieu, Pierre")
