"""تصدير النصوص إلى Word بالهوية البصرية (اتجاه من اليمين، Noto Naskh Arabic، عناوين بلون الإدارة وخط ذهبي).

clean=True نسخة القراءة: تُحذف وسوم التصنيف، وتُترجم مواضع النقص والقرار إلى العربية بين قوسين.
clean=False نسخة المراجعة: كل الوسوم بالعربية بين قوسين. وفي الحالين تُنقّى العلامات المائية الظاهرة والمخفية
(runtime/rkpos/sanitize.py)، والأصل (Markdown) يبقى بوسومه المعيارية.
"""
from __future__ import annotations
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import sanitize

BLUE, GOLD = RGBColor(0x1F, 0x4E, 0x79), RGBColor(0xB8, 0x86, 0x0B)
FONT = "Noto Naskh Arabic"


def _rtl(par, align="both"):
    ppr = par._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.append(bidi)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), {"start": "both", "center": "center", "both": "both"}[align])
    ppr.append(jc)


def _run(par, text, size=14, bold=False, color=None):
    """مقطع نصي بخط عربي واتجاه من اليمين؛ تُدرج عناصر الخصائص في مواضعها من المخطط (يرفض Word الترتيب الخاطئ)."""
    r = par.add_run(text)
    r.font.name, r.font.size, r.bold = FONT, Pt(size), bold
    if color:
        r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), FONT)
    if bold:
        rpr.find(qn("w:b")).addnext(OxmlElement("w:bCs"))
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(size * 2))
    rpr.find(qn("w:sz")).addnext(szcs)
    rpr.append(OxmlElement("w:rtl"))
    return r


def _inline(par, text, size):
    """**غامق** داخل الفقرة."""
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            _run(par, chunk, size=size, bold=bool(i % 2))


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def to_docx(markdown: str, title: str | None = None, clean: bool = True, theme: dict | None = None,
            author: str = "") -> bytes:
    """theme: ألوان الإدارة من نظام التصميم المركزي (primary للعناوين، accent للخط الفاصل)."""
    blue = _rgb(theme["primary"]) if theme else BLUE
    gold = _rgb(theme["accent"]) if theme else GOLD
    markdown, _ = sanitize.clean(markdown, "reading" if clean else "tagged")
    if title:   # لا يتكرر العنوان إن كان النص يبدأ به
        markdown = re.sub(rf"\A#\s+{re.escape(title.strip())}\s*\n", "", markdown.lstrip())
    doc = Document()
    sec = doc.sections[0]
    sec.right_to_left = True
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(14)
    if title:
        p = doc.add_paragraph()
        _rtl(p, "center")
        _run(p, title, size=24, bold=True, color=blue)
        line = doc.add_paragraph()
        _rtl(line, "center")
        _run(line, "━━━━━━━━", size=12, color=gold)
    for block in re.split(r"\n\s*\n", markdown.strip()):
        block = block.strip()
        if not block or block.startswith(("```", "===")):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", block)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            _rtl(p, "start")
            p.paragraph_format.space_before = Pt(14)
            _run(p, m.group(2).strip(), size={1: 20, 2: 17, 3: 15}.get(level, 14), bold=True, color=blue)
            continue
        for line in block.split("\n"):
            line = line.rstrip()
            if not line:
                continue
            bullet = re.match(r"^\s*[-*•]\s+(.*)", line)
            p = doc.add_paragraph()
            _rtl(p, "both")
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.4
            _inline(p, ("• " + bullet.group(1)) if bullet else line, 14)
    sanitize.scrub_docx(doc, author=author, title=title or "")
    buf = io.BytesIO()
    doc.save(buf)
    return sanitize.scrub_package(buf.getvalue())
